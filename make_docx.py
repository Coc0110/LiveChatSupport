# -*- coding: utf-8 -*-
"""
Script tao file Bao cao Do an Mon hoc: CMP180 - Lap trinh mang may tinh
De tai: HE THONG CHAT HO TRO KHACH HANG TRUC TUYEN (LIVE CHAT SUPPORT)
GVHD: ThS. Nguyen Van Can
Sinh vien: Hoang Trung Anh (2380602692), Huynh Phat Loc (2380613189)
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_report():
    doc = Document()

    # 1. Page Margins (A4: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)      # 2.0 cm
        section.bottom_margin = Inches(0.79)   # 2.0 cm
        section.left_margin = Inches(1.18)     # 3.0 cm (đóng gáy sách)
        section.right_margin = Inches(0.79)    # 2.0 cm
        section.page_width = Inches(8.27)      # A4
        section.page_height = Inches(11.69)

    # 2. Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    style_normal.paragraph_format.line_spacing = 1.3
    style_normal.paragraph_format.space_after = Pt(6)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Helper styling functions
    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                          f'<w:top w:w="{top}" w:type="dxa"/>'
                          f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                          f'<w:left w:w="{left}" w:type="dxa"/>'
                          f'<w:right w:w="{right}" w:type="dxa"/>'
                          f'</w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                               f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                               f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                               f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                               f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                               f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                               f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                               f'</w:tblBorders>')
        tblPr.append(tblBorders)

    def add_p(text, bold_prefix=None, italic=False, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = align
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(13)
        r = p.add_run(text)
        r.italic = italic
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        return p

    def add_heading_chapter(title):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78) # Navy Blue
        return p

    def add_heading_1(title):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        return p

    def add_heading_2(title):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.bold = True
        r.italic = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        return p

    def add_heading_3(title):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_code_block(code_str, caption=None):
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(6)
            p_cap.paragraph_format.space_after = Pt(2)
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_c = p_cap.add_run(caption)
            r_c.font.name = "Times New Roman"
            r_c.font.size = Pt(10.5)
            r_c.italic = True
            r_c.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                            f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0078D7"/>'
                            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'</w:tcBorders>')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        r = p.add_run(code_str.strip())
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(4)

    def add_table_custom(headers, rows_data, col_widths=None, caption=None):
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_c = p_cap.add_run(caption)
            r_c.bold = True
            r_c.font.name = "Times New Roman"
            r_c.font.size = Pt(11)
            r_c.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="B0C4DE", sz="4", val="single")

        # Header Row
        hdr_cells = tbl.rows[0].cells
        for idx, header_text in enumerate(headers):
            cell = hdr_cells[idx]
            cell.text = header_text
            set_cell_background(cell, "1F4E78") # Navy header
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Body Rows
        for r_idx, row_values in enumerate(rows_data):
            row_cells = tbl.rows[r_idx + 1].cells
            bg_color = "F4F8FC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row_cells[c_idx]
                cell.text = str(val)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 1 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        # Apply Column Widths if provided
        if col_widths:
            for row in tbl.rows:
                for c_idx, width in enumerate(col_widths):
                    row.cells[c_idx].width = Inches(width)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)

    print("Setup complete. Generating sections...")
    return doc, {
        'add_p': add_p,
        'add_heading_chapter': add_heading_chapter,
        'add_heading_1': add_heading_1,
        'add_heading_2': add_heading_2,
        'add_heading_3': add_heading_3,
        'add_code_block': add_code_block,
        'add_table_custom': add_table_custom
    }

print("make_docx base ready.")

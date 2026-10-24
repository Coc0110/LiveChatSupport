# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                           f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                           f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                           f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                           f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                           f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                           f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                           f'</w:tblBorders>')
    tblPr.append(tblBorders)

def add_callout(doc, text, title=None, hex_bg="EBF3FA", hex_border="2E75B6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=140, bottom=140, left=220, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{hex_border}"/>'
                        f'<w:top w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:bottom w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(2)
    if title:
        r_title = p.add_run(f"★ {title}\n")
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(11.5)
        r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    
    r_body = p.add_run(text)
    r_body.italic = True
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_c = p_cap.add_run(caption)
        r_c.font.name = "Times New Roman"
        r_c.font.size = Pt(10.5)
        r_c.italic = True
        r_c.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0078D7"/>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    
    r = p.add_run(code_str.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

print("Helper definitions loaded.")
